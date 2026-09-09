"""Excel output adapter.

Transcribes `report_manifest.json` sections into worksheets — one sheet per
section, `bar_chart` sections also get an openpyxl `BarChart` object. This
module holds no aggregation or judgement logic of its own; every number
already came from Phase 2/3.
"""

from __future__ import annotations

import argparse
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from core import db as db_module
from output.report_manifest import ReportManifestBuilder, build_manifest


def _safe_sheet_title(title: str, used_titles: set[str]) -> str:
    """Excel sheet titles are capped at 31 chars and must be unique per workbook."""
    base = re.sub(r"[\[\]:*?/\\]", "_", title)[:31]
    candidate = base
    suffix_n = 2
    while candidate in used_titles:
        suffix = f"_{suffix_n}"
        candidate = base[: 31 - len(suffix)] + suffix
        suffix_n += 1
    used_titles.add(candidate)
    return candidate


def write_workbook(builder: ReportManifestBuilder, out_path: str) -> None:
    manifest = builder.build()
    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()

    for section in manifest["sections"]:
        ws = wb.create_sheet(
            _safe_sheet_title(section["title"] or section["section_id"], used_titles)
        )
        data = section["data"]

        if section["type"] == "bar_chart":
            ws.append(["label", "value"])
            for label, value in zip(data["labels"], data["values"]):
                ws.append([label, value])

            chart = BarChart()
            chart.title = section["title"]
            n_rows = len(data["labels"])
            values_ref = Reference(ws, min_col=2, min_row=1, max_row=n_rows + 1)
            cats_ref = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
            chart.add_data(values_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws.add_chart(chart, "D2")

        elif section["type"] == "table":
            ws.append(data["columns"])
            for row in data["rows"]:
                ws.append(row)

        elif section["type"] == "diff_summary":
            ws.append(["previous_scan_run_id", data["previous_scan_run_id"]])
            ws.append(["added_count", data["added_count"]])
            ws.append(["changed_count", data["changed_count"]])
            ws.append(["removed_count", data["removed_count"]])
            ws.append([])
            for kind in ("added_ids", "changed_ids", "removed_ids"):
                ws.append([kind])
                for asset_id in data[kind]:
                    ws.append(["", asset_id])

        elif section["type"] == "dot_chart":
            # SVG has no Excel equivalent; transcribe the same per-point data
            # (is_outlier already decided by trend/aggregations.py) as a table.
            ws.append(["mean", data.get("mean")])
            ws.append(["outlier_method", data.get("outlier_method")])
            ws.append([])
            ws.append(["asset_id", "value", "is_outlier"])
            for point in data.get("points", []):
                ws.append([point["asset_id"], point["value"], point["is_outlier"]])

    wb.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build the Excel report from report_manifest data"
    )
    parser.add_argument("--db", default="insight.db")
    parser.add_argument("--scan-run-id", type=int, default=None)
    parser.add_argument("--out", default="report.xlsx")
    args = parser.parse_args()

    with db_module.open_db(args.db) as conn:
        scan_run_id = args.scan_run_id or db_module.latest_scan_run_id(conn)
        if scan_run_id is None:
            raise SystemExit("no successful scan_run found")
        builder = build_manifest(conn, scan_run_id)

    write_workbook(builder, args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
